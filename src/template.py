"""Builds the Excel input workbook (also used to export any company back to Excel)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import CompanyData
from .samples import example_company
from .schema import BALANCE_INPUTS, INCOME_INPUTS, LABELS, NOTES

FONT = "Arial"
NAVY = "1F3864"
NUM_FMT = '#,##0.0;(#,##0.0);"-"'

# ("input", key) | ("calc", key, formula using {c}=column letter and {other_key}=row number) | ("section", title)
INCOME_LAYOUT = [
    ("input", "revenue"), ("input", "cogs"),
    ("calc", "gross_profit", "={c}{revenue}-{c}{cogs}"),
    ("input", "sga"), ("input", "other_opex"),
    ("calc", "ebit", "={c}{gross_profit}-{c}{sga}-{c}{other_opex}"),
    ("input", "interest_expense"), ("input", "other_income"),
    ("calc", "pretax", "={c}{ebit}-{c}{interest_expense}+{c}{other_income}"),
    ("input", "tax"),
    ("calc", "net_income", "={c}{pretax}-{c}{tax}"),
    ("section", "Memo items"),
    ("input", "da"), ("input", "dividends"), ("input", "shares"),
]
BALANCE_LAYOUT = [
    ("section", "Assets"),
    ("input", "cash"), ("input", "st_investments"), ("input", "receivables"),
    ("input", "inventory"), ("input", "other_current_assets"),
    ("calc", "total_current_assets", "=SUM({c}{cash}:{c}{other_current_assets})"),
    ("input", "ppe"), ("input", "intangibles"), ("input", "other_noncurrent_assets"),
    ("calc", "total_assets", "={c}{total_current_assets}+SUM({c}{ppe}:{c}{other_noncurrent_assets})"),
    ("section", "Liabilities & Equity"),
    ("input", "payables"), ("input", "st_debt"), ("input", "accrued_liabilities"),
    ("calc", "total_current_liabilities", "=SUM({c}{payables}:{c}{accrued_liabilities})"),
    ("input", "lt_debt"), ("input", "other_noncurrent_liabilities"),
    ("calc", "total_liabilities", "={c}{total_current_liabilities}+{c}{lt_debt}+{c}{other_noncurrent_liabilities}"),
    ("input", "share_capital"), ("input", "retained_earnings"),
    ("calc", "total_equity", "={c}{share_capital}+{c}{retained_earnings}"),
    ("calc", "total_liab_equity", "={c}{total_liabilities}+{c}{total_equity}"),
    ("calc", "balance_check", "={c}{total_assets}-{c}{total_liab_equity}"),
]


def _statement_sheet(ws, layout, df):
    years = list(df.columns)
    notes_col = len(years) + 2
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1, "Line Item")
    for j, y in enumerate(years, start=2):
        ws.cell(1, j, str(y))  # years stored as text
    ws.cell(1, notes_col, "Notes")
    for col in range(1, notes_col + 1):
        cell = ws.cell(1, col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="left" if col in (1, notes_col) else "right")

    rows, r = {}, 2
    for item in layout:
        if item[0] != "section":
            rows[item[1]] = r
        r += 1

    r = 2
    for item in layout:
        kind = item[0]
        if kind == "section":
            ws.cell(r, 1, item[1]).font = Font(name=FONT, bold=True, color=NAVY)
        else:
            key = item[1]
            ws.cell(r, 1, LABELS[key]).font = Font(name=FONT, bold=(kind == "calc"))
            for j, y in enumerate(years, start=2):
                cell = ws.cell(r, j)
                cell.number_format = NUM_FMT
                if kind == "input":
                    cell.value = float(df.loc[key, y])
                    cell.font = Font(name=FONT, color="0000FF")  # blue = input
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                else:
                    cell.value = item[2].format(c=get_column_letter(j), **rows)
                    cell.font = Font(name=FONT, bold=True)
                    cell.border = Border(top=Side(style="thin"))
            if key in NOTES:
                ws.cell(r, notes_col, NOTES[key]).font = Font(name=FONT, italic=True, color="7F7F7F")
        r += 1
    ws.column_dimensions["A"].width = 40
    for j in range(2, notes_col):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.column_dimensions[get_column_letter(notes_col)].width = 55
    ws.freeze_panes = "B2"


def build_workbook(c: CompanyData) -> bytes:
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    info["A1"] = "Company Information"
    info["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    fields = [("Company Name", c.name), ("Ticker", c.ticker), ("Currency", c.currency),
              ("Units", c.unit), ("Sector", c.sector)]
    for i, (k, v) in enumerate(fields, start=3):
        info.cell(i, 1, k).font = Font(name=FONT, bold=True)
        cell = info.cell(i, 2, v)
        cell.font = Font(name=FONT, color="0000FF")
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    legend = [
        "How to use this file",
        "1. Edit only the blue cells (yellow background). Black rows are calculated and ignored by the program.",
        "2. Enter all expenses as POSITIVE numbers (the program subtracts them).",
        "3. Row 1 of each sheet holds the fiscal years. Add or remove year columns as needed (4-digit year header).",
        "4. Do not rename the line items in column A, and keep the sheet names: Info, IncomeStatement, BalanceSheet.",
        "5. Depreciation & Amortization is a memo item: it is already inside COGS / SG&A and is only used for EBITDA.",
        "6. The values in this template are illustrative (fictional) - replace them with your company's figures.",
    ]
    for i, t in enumerate(legend, start=10):
        info.cell(i, 1, t).font = Font(name=FONT, bold=(i == 10), color=NAVY if i == 10 else "000000")
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 45

    _statement_sheet(wb.create_sheet("IncomeStatement"), INCOME_LAYOUT, c.income)
    _statement_sheet(wb.create_sheet("BalanceSheet"), BALANCE_LAYOUT, c.balance)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_bytes() -> bytes:
    return build_workbook(example_company())
